"""DEV-ONLY: builds the synthetic Marketing report workbooks used by the AURA end-to-end test.

Same layout as the real AM report (row 1 title, row 2 summary, row 3 blank, row 4 headers).
Every account, person and address is fictitious (example.test). No customer data.
Usage: python tools/aura-e2e/make_synthetic_report.py
"""
import os
from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "fixtures")
HEADERS = ["Cuenta", "Account Owner", "Agente responsable (Sales Rep Actual)", "Pais Billing", "Pais Shipping",
           "Contacto", "Email", "Posicion (Company position)", "Titulo", "Key Contact", "Prioridad", "Motivo campana"]


def row(account, owner, country, contact, email, priority, reason):
    return [account, owner, owner, country, country, contact, email, "Operaciones", "Gerente", "Si", priority, reason]


def tab(wb, name, title, rows):
    ws = wb.create_sheet(name)
    ws.append([title])
    ws.append(["09/10/2026 | %d contactos | Prioridad: A" % len(rows)])
    ws.append([])
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)


def report(version):
    wb = Workbook()
    wb.remove(wb.active)
    tab(wb, "Retencion prioritaria", "DGL Freight Broker - Retencion prioritaria (QA)", [
        row("QA Transportes Andinos", "AM Uno QA", "Colombia", "Ana QA", "ana@andinos.example.test", "Alta", "Cuenta en riesgo (bajo de tier)"),
        row("QA Logistica Pacifico", "AM Dos QA", "Peru", "Luis QA", "luis@pacifico.example.test", "Media", "Volumen en baja este mes"),
        row("QA Cargo Brasil", "AM Uno QA", "Brazil", "Joao QA", "joao@cargobr.example.test", "Alta", "Cliente clave sin cotizaciones recientes"),
        row("", "AM Uno QA", "Mexico", "Sin Cuenta QA", "x@sincuenta.example.test", "Baja", "Fila invalida: sin cuenta"),
        row("QA Sin Responsable", "", "Mexico", "Sin AM QA", "y@sinam.example.test", "Baja", "Fila invalida: sin responsable"),
    ])
    tab(wb, "Recuperacion FTL", "DGL Freight Broker - Recuperacion FTL (QA)", [
        row("QA Freight Texas", "AM Tres QA", "United States", "Mike QA", "mike@freighttx.example.test", "Alta", "Area FTL en baja"),
        row("QA Transportes Andinos", "AM Uno QA", "Colombia", "Ana QA", "ana@andinos.example.test", "Media", "Tambien en recuperacion FTL"),
        row("QA Mudanzas Norte", "AM Tres QA", "Mexico", "Rosa QA", "rosa@mudanzas.example.test", "Media", "Sin cargas en 60 dias"),
    ])
    expansion = [row("QA Distribuidora Sur", "AM Dos QA", "Chile", "Pedro QA", "pedro@dsur.example.test", "Media", "Usa un solo servicio")]
    if version == 2:
        expansion.append(row("QA Agro Export", "AM Dos QA", "Peru", "Carla QA", "carla@agro.example.test", "Alta", "Nuevo interes en cross border"))
    tab(wb, "Expansion de servicio", "DGL Freight Broker - Expansion de servicio (QA)", expansion)
    tab(wb, "Confirmar datos contacto", "DGL Freight Broker - Confirmar datos de contacto (QA)", [
        row("QA Datos Pendientes", "AM Uno QA", "Mexico", "Dato QA", "", "Alta", "Falta email"),
        row("QA Datos Pendientes 2", "AM Dos QA", "Peru", "Dato2 QA", "", "Alta", "Falta telefono"),
    ])
    tab(wb, "Campana B - nueva", "Campana B (pestana no reconocida)", [
        row("QA Campana B", "AM Tres QA", "Mexico", "B QA", "b@campb.example.test", "Alta", "Pestana nueva sin mapeo"),
    ])
    tab(wb, "Campana A - HA prioritaria", "Campana A (pipeline propio, cerrado)", [
        row("QA House Account", "AM Uno QA", "Mexico", "HA QA", "ha@house.example.test", "Alta", "House account"),
    ])
    path = os.path.join(OUT, "Marketing_DGL_QA_SINTETICO_v%d.xlsx" % version)
    wb.save(path)
    return path


def not_a_report():
    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto Q4"
    ws.append(["Concepto", "Monto"])
    ws.append(["Viajes QA", 1000])
    path = os.path.join(OUT, "Presupuesto_Q4_QA.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for p in (report(1), report(2), not_a_report()):
        print(p)
