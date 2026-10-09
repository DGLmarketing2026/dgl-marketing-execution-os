"""DEV-ONLY stand-in for the Drive XLSX -> Google Sheets conversion used by AURA.

Reads a real .xlsx (path in argv[1]) and prints, per tab, the 2-D values that
Sheet.getDataRange().getValues() would return (empty cells as "").
"""
import json
import sys
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1], data_only=True)
out = []
for ws in wb.worksheets:
    values = [["" if v is None else v for v in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True)]
    out.append({"name": ws.title, "values": values})
print(json.dumps(out, ensure_ascii=False))
